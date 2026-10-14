"""Converte a planilha de convidados em SQL para o banco.

Uso:
    python scripts/importar-convidados.py assets/convidados.xlsx > assets/convidados.sql

A planilha tem duas abas: "Convidados" (código G###, nome, telefone, e-mail) e
"Acompanhantes" (código A###, nome e até dois códigos de convidado principal).
O SQL gerado é idempotente: rodar de novo atualiza nomes e contatos pelo código
sem duplicar ninguém e sem apagar respostas já dadas.

A saída contém dados pessoais e por isso vai para assets/, que fica fora do git.
"""

import sys

from openpyxl import load_workbook

HEADER_ROW = 4  # as três primeiras linhas são título e instruções


def sql(value):
    if value is None or str(value).strip() == "":
        return "null"
    return "'" + str(value).strip().replace("'", "''") + "'"


def text(cell):
    # O Excel guarda telefone como número: 11981017458 viria como "11981017458.0".
    if isinstance(cell, float) and cell.is_integer():
        cell = int(cell)
    return None if cell is None else str(cell).strip()


def rows(sheet):
    for row in sheet.iter_rows(min_row=HEADER_ROW + 1, values_only=True):
        if row and row[0] and row[1] and str(row[1]).strip():
            yield [text(cell) for cell in row]


def main(path):
    book = load_workbook(path, read_only=True, data_only=True)
    principals = list(rows(book["Convidados"]))
    companions = list(rows(book["Acompanhantes"]))

    known = {code for code, *_ in principals}
    out = ["begin;", ""]

    for code, name, phone, email, *_ in principals:
        out.append(
            "insert into public.guests (code, full_name, kind, phone, email) "
            f"values ({sql(code)}, {sql(name)}, 'principal', {sql(phone)}, {sql(email)}) "
            "on conflict (code) do update set full_name = excluded.full_name, "
            "phone = excluded.phone, email = excluded.email;"
        )

    out.append("")
    orphans = []

    for code, name, first, second, *_ in companions:
        out.append(
            "insert into public.guests (code, full_name, kind) "
            f"values ({sql(code)}, {sql(name)}, 'companion') "
            "on conflict (code) do update set full_name = excluded.full_name;"
        )
        links = [link for link in (first, second) if link]
        if not links or any(link not in known for link in links):
            orphans.append(f"{code} {name} -> {links}")
        for link in links:
            out.append(
                "insert into public.guest_links (principal_id, companion_id) "
                f"select p.id, c.id from public.guests p, public.guests c "
                f"where p.code = {sql(link)} and c.code = {sql(code)} on conflict do nothing;"
            )

    out += ["", "commit;"]
    print("\n".join(out))

    # Vai para o stderr: não entra no SQL, mas aparece no terminal.
    print(f"-- {len(principals)} convidados, {len(companions)} acompanhantes", file=sys.stderr)
    for orphan in orphans:
        print(f"-- ATENÇÃO, vínculo inexistente: {orphan}", file=sys.stderr)


if __name__ == "__main__":
    main(sys.argv[1])
